import io
import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from xhtml2pdf import pisa
from django.http import HttpResponse

# Caracteres de control que openpyxl/XML 1.0 no admite en celdas
ILLEGAL_XML_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]')


def sanitize_cell(value):
    """
    Limpia un valor para que pueda escribirse en una celda de Excel.
    Elimina caracteres de control ilegales (ej: \x02) que openpyxl rechaza.
    """
    if isinstance(value, str):
        return ILLEGAL_XML_CHARS.sub('', value)
    return value

def export_report_pdf(html_string, filename="reporte_contable.pdf"):
    """
    Convierte una plantilla HTML renderizada a un archivo PDF listo para descarga.
    """
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html_string.encode("UTF-8")), result)
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    return HttpResponse("Error generando el archivo PDF", status=500)


def export_trial_balance_excel(company, trial_data):
    """
    Genera un libro de Excel (.xlsx) para el Balance de Comprobación.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance de Comprobación"

    # Estilos
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True)
    subtitle_font = Font(name="Arial", size=10, italic=True)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append([sanitize_cell(company.name)])
    ws.append(["BALANCE DE COMPROBACIÓN (SUMAS Y SALDOS)"])
    ws.append([f"RUC: {sanitize_cell(company.tax_id)}"])
    ws.append([])

    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=2, column=1).font = Font(name="Arial", size=12, bold=True)
    ws.cell(row=3, column=1).font = subtitle_font

    headers = ["Código", "Cuenta Contable", "Suma Debe", "Suma Haber", "Saldo Deudor", "Saldo Acreedor"]
    ws.append(headers)

    header_row = 5
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for line in trial_data['lines']:
        ws.append([
            sanitize_cell(line['account'].code),
            sanitize_cell(line['account'].name),
            float(line['sum_debit']),
            float(line['sum_credit']),
            float(line['debt_balance']),
            float(line['cred_balance']),
        ])

    total_row = [
        "TOTALES",
        "",
        float(trial_data['total_debit_sum']),
        float(trial_data['total_credit_sum']),
        float(trial_data['total_debt_balance']),
        float(trial_data['total_cred_balance']),
    ]
    ws.append(total_row)

    last_row = ws.max_row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=last_row, column=col_num)
        cell.font = Font(name="Arial", size=11, bold=True)
        cell.border = Border(top=Side(style='thin'), bottom=Side(style='double'))

    # Ajustar anchos de columna
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Balance_Comprobacion_{company.tax_id}.xlsx"'
    return response


def _excel_styles():
    """Estilos compartidos para las exportaciones de Excel."""
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True)
    subtitle_font = Font(name="Arial", size=10, italic=True)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    return {
        'header_font': header_font,
        'title_font': title_font,
        'subtitle_font': subtitle_font,
        'header_fill': header_fill,
        'thin_border': thin_border,
    }


def export_journal_excel(company, entries, start_date, end_date, selected_account_ids=None):
    """
    Genera un libro de Excel (.xlsx) profesional para el Libro Diario,
    con una fila por línea de asiento (código, cuenta, detalle y referencia)
    y subtotales por asiento.
    """
    selected_ids = set(selected_account_ids) if selected_account_ids else None
    styles = _excel_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libro Diario"

    ws.append([sanitize_cell(company.name)])
    ws.append(["LIBRO DIARIO (ASIENTOS CONTABLES)"])
    ws.append([f"RUC: {sanitize_cell(company.tax_id)} | Periodo: {start_date} al {end_date}"])
    ws.append([])

    ws.cell(row=1, column=1).font = styles['title_font']
    ws.cell(row=2, column=1).font = Font(name="Arial", size=12, bold=True)
    ws.cell(row=3, column=1).font = styles['subtitle_font']

    headers = ["Nro", "Fecha", "Tipo", "Referencia", "Código", "Cuenta Contable", "Detalle", "Debe ($)", "Haber ($)"]
    ws.append(headers)

    header_row = 5
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row = header_row + 1
    total_debit = 0.0
    total_credit = 0.0
    bold = Font(name="Arial", size=10, bold=True)

    for entry in entries:
        sum_d = 0.0
        sum_c = 0.0
        for line in entry.lines.all():
            if selected_ids and line.account_id not in selected_ids:
                continue
            ws.cell(row=row, column=1, value=entry.entry_number)
            ws.cell(row=row, column=2, value=entry.date.isoformat())
            ws.cell(row=row, column=3, value=entry.get_entry_type_display())
            ws.cell(row=row, column=4, value=sanitize_cell(entry.reference or ""))
            ws.cell(row=row, column=5, value=sanitize_cell(line.account.code))
            ws.cell(row=row, column=6, value=sanitize_cell(line.account.name))
            ws.cell(row=row, column=7, value=sanitize_cell(line.concept or ""))
            ws.cell(row=row, column=8, value=float(line.debit) if line.debit else None)
            ws.cell(row=row, column=9, value=float(line.credit) if line.credit else None)
            row += 1
            sum_d += float(line.debit) if line.debit else 0.0
            sum_c += float(line.credit) if line.credit else 0.0

        # Subtotal del asiento
        subtotal_label = f"SUBTOTAL ASIENTO #{entry.entry_number}" if not selected_ids else "SUBTOTAL"
        ws.cell(row=row, column=4, value=subtotal_label).font = bold
        ws.cell(row=row, column=8, value=sum_d).font = bold
        ws.cell(row=row, column=9, value=sum_c).font = bold
        for col_num in range(1, len(headers) + 1):
            ws.cell(row=row, column=col_num).border = Border(top=Side(style='thin'))
        row += 1
        total_debit += sum_d
        total_credit += sum_c

    ws.cell(row=row, column=4, value="TOTALES GENERALES").font = Font(name="Arial", size=11, bold=True)
    ws.cell(row=row, column=8, value=total_debit).font = Font(name="Arial", size=11, bold=True)
    ws.cell(row=row, column=9, value=total_credit).font = Font(name="Arial", size=11, bold=True)
    for col_num in range(1, len(headers) + 1):
        ws.cell(row=row, column=col_num).border = Border(top=Side(style='thin'), bottom=Side(style='double'))

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 13
    ws.column_dimensions['F'].width = 40
    ws.column_dimensions['G'].width = 40
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['I'].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Libro_Diario_{company.tax_id}.xlsx"'
    return response


def export_ledger_excel(company, data):
    """
    Genera un libro de Excel (.xlsx) profesional para el Libro Mayor,
    con una sección por cuenta (saldo inicial, movimientos, sumas y saldo final).
    """
    styles = _excel_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libro Mayor"

    ws.append([sanitize_cell(company.name)])
    ws.append(["LIBRO MAYOR"])
    ws.append([f"RUC: {sanitize_cell(company.tax_id)} | Periodo: {data['start_date']} al {data['end_date']}"])
    ws.append([])

    ws.cell(row=1, column=1).font = styles['title_font']
    ws.cell(row=2, column=1).font = Font(name="Arial", size=12, bold=True)
    ws.cell(row=3, column=1).font = styles['subtitle_font']

    row = 5
    headers = ["Fecha", "Detalle", "Referencia", "Debe ($)", "Haber ($)", "Saldo ($)"]

    for item in data['accounts']:
        # Encabezado de cuenta
        ws.cell(row=row, column=1, value=sanitize_cell(f"{item['account'].code} - {item['account'].name}")).font = Font(name="Arial", size=11, bold=True)
        row += 1

        # Fila de cabecera de columnas
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_num)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row += 1

        # Saldo inicial
        ws.cell(row=row, column=1, value="SALDO INICIAL")
        ws.cell(row=row, column=6, value=float(item['opening_balance'])).font = Font(name="Arial", size=10, italic=True)
        row += 1

        # Movimientos
        for m in item['movements']:
            ws.cell(row=row, column=1, value=m['date'].isoformat())
            ws.cell(row=row, column=2, value=sanitize_cell(f"#{m['entry_number']} {m['concept']}"))
            ws.cell(row=row, column=3, value=sanitize_cell(m['reference']))
            ws.cell(row=row, column=4, value=float(m['debit']) if m['debit'] else None)
            ws.cell(row=row, column=5, value=float(m['credit']) if m['credit'] else None)
            ws.cell(row=row, column=6, value=float(m['balance']))
            row += 1

        # Sumas y saldo final
        ws.cell(row=row, column=2, value="SUMAS").font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=row, column=4, value=float(item['total_debit']))
        ws.cell(row=row, column=5, value=float(item['total_credit']))
        ws.cell(row=row, column=6, value=float(item['closing_balance'])).font = Font(name="Arial", size=10, bold=True)
        row += 2

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 16
    ws.column_dimensions['F'].width = 16

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Libro_Mayor_{company.tax_id}.xlsx"'
    return response
